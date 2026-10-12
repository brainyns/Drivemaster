"""
Generador de Documento Word (.docx)
Explicación detallada del Modelo Predictivo de Aprobación de Cotizaciones - DriveMaster
Autores: Brainy Blanco, Sebastian Ortega, Juan Cantillo, Elias Puello
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def crear_documento():
    doc = Document()
    
    # Configuración de márgenes
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    COLOR_PRIMARIO = RGBColor(15, 76, 129)     # Azul Institucional
    COLOR_SECUNDARIO = RGBColor(30, 41, 59)   # Slate oscuro
    COLOR_GRIS = RGBColor(100, 116, 139)       # Gris texto secundario
    COLOR_VERDE = RGBColor(16, 185, 129)      # Verde acento
    COLOR_FONDO_TABLA = "F1F5F9"
    COLOR_HEADER_TABLA = "0F4C81"

    # --- PORTADA / ENCABEZADO ---
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(10)
    title_p.paragraph_format.space_after = Pt(4)
    run_sub = title_p.add_run("PROPUESTA E IMPLEMENTACIÓN · PROYECTO DE AULA")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(11)
    run_sub.font.bold = True
    run_sub.font.color.rgb = COLOR_PRIMARIO

    h1_p = doc.add_paragraph()
    h1_p.paragraph_format.space_after = Pt(14)
    run_title = h1_p.add_run("Modelo Predictivo para Estimar la Probabilidad de Aprobación de una Cotización de Servicio en DriveMaster")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_SECUNDARIO

    # Info card
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(20)
    r1 = p_meta.add_run("Integrantes: ")
    r1.bold = True
    p_meta.add_run("Brainy Blanco · Sebastian Ortega · Juan Cantillo · Elias Puello\n")
    r2 = p_meta.add_run("Plataforma: ")
    r2.bold = True
    p_meta.add_run("DriveMaster (Marketplace Automotriz & Red de Talleres)\n")
    r3 = p_meta.add_run("Módulo: ")
    r3.bold = True
    p_meta.add_run("Inteligencia de Negocio y Minería de Datos (Machine Learning)")

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # --- SECCIÓN 1: INTRODUCCIÓN Y PROBLEMA ---
    h_sec1 = doc.add_heading("1. Descripción del Problema y Justificación", level=1)
    h_sec1.style.font.color.rgb = COLOR_PRIMARIO

    p1 = doc.add_paragraph(
        "En la plataforma DriveMaster, los talleres mecánicos reciben solicitudes de diagnóstico y mantenimiento "
        "de clientes automotrices. Para responder a una solicitud, el taller formula una cotización formal que incluye "
        "los repuestos (provenientes de tiendas de autopartes asociadas), la mano de obra calificada y el IVA del 19 %."
    )
    p2 = doc.add_paragraph(
        "Tradicionalmente, el taller no cuenta con ninguna herramienta para anticipar si el cliente va a aceptar "
        "o a rechazar la propuesta. Cada cotización rechazada genera consecuencias negativas para el ecosistema:"
    )
    
    b1 = doc.add_paragraph(style='List Bullet')
    r = b1.add_run("Tiempo invertido no remunerado: ")
    r.bold = True
    b1.add_run("El personal del taller dedica tiempo inspeccionando y cotizando sin concretar el trabajo.")

    b2 = doc.add_paragraph(style='List Bullet')
    r = b2.add_run("Pérdida de ventas en tiendas asociadas: ")
    r.bold = True
    b2.add_run("Los repuestos cotizados que no se compran representan lucro cesante para los comercios.")

    b3 = doc.add_paragraph(style='List Bullet')
    r = b3.add_run("Fuga del cliente: ")
    r.bold = True
    b3.add_run("Un cliente que percibe un precio desproporcionado o sufre una respuesta tardía abandona la plataforma.")

    doc.add_paragraph(
        "La integración del modelo predictivo resuelve este problema suministrando una alerta temprana en tiempo real "
        "al taller en el preciso momento en que está estructurando la cotización, permitiéndole tomar decisiones correctivas "
        "(como ajustar tarifas, cambiar repuestos o priorizar clientes) antes de enviarla."
    )

    # --- SECCIÓN 2: OBJETIVO Y VARIABLE OBJETIVO ---
    h_sec2 = doc.add_heading("2. Definición del Modelo y Variable Objetivo", level=1)
    h_sec2.style.font.color.rgb = COLOR_PRIMARIO

    doc.add_paragraph(
        "El objetivo es desarrollar un modelo predictivo de clasificación supervisada capaz de estimar "
        "la probabilidad matemática (entre 0.0 y 1.0) de que una cotización sea aprobada por el cliente."
    )

    p_var = doc.add_paragraph()
    r = p_var.add_run("Variable Objetivo (Target): ")
    r.bold = True
    p_var.add_run("aprobacion_cotizacion\n")
    r = p_var.add_run("Tipo: ")
    r.bold = True
    p_var.add_run("Categórica Binaria (1 = Sí / Aprobada, 0 = No / Rechazada o Cancelada).\n")
    r = p_var.add_run("Origen en el Sistema: ")
    r.bold = True
    p_var.add_run("Se mapea del estado final de las solicitudes y órdenes en la base de datos (estados APROBADA vs. RECHAZADA/CANCELADA).")

    # --- SECCIÓN 3: CRITERIOS Y VARIABLES PREDICTORAS ---
    h_sec3 = doc.add_heading("3. Criterios de Predicción y Variables Predictoras (Explicación Detallada)", level=1)
    h_sec3.style.font.color.rgb = COLOR_PRIMARIO

    doc.add_paragraph(
        "El modelo evalúa 12 variables predictoras extraídas del contexto del vehículo, el taller, la logística "
        "y los componentes económicos de la cotización. A continuación se detalla cada criterio y su impacto en la decisión del cliente:"
    )

    # Tabla de variables
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    headers = ["Variable", "Tipo de Dato", "Significado en la App", "Impacto en la Probabilidad"]
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_background(hdr_cells[i], COLOR_HEADER_TABLA)
        set_cell_margins(hdr_cells[i], 140, 140, 180, 180)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)

    criterios_data = [
        ("tiempo_respuesta_horas", "Numérica continua", "Horas transcurridas desde que el cliente creó la solicitud hasta que el taller formuló la cotización.", "CRÍTICO (-0.647): Mayor demora disminuye drásticamente la probabilidad. Responder en las primeras 4h eleva la aprobación."),
        ("calificacion_taller", "Numérica (1.0 - 5.0)", "Reputación y rating promedio obtenido por el taller a partir de reseñas de clientes previos.", "POSITIVO (+0.524): Talleres con rating ≥ 4.6 generan mayor confianza e inducen mayor tasa de aceptación."),
        ("distancia_km", "Numérica continua", "Distancia geográfica estimada entre la ubicación del cliente y la sede física del taller.", "NEGATIVO (-0.400): Distancias mayores a 15-20 km desalientan al cliente por costos de desplazamiento y grúa."),
        ("anio_vehiculo", "Numérica entera", "Año del vehículo (antigüedad).", "POSITIVO (+0.312): Autos más recientes suelen tener dueños dispuestos a realizar mantenimientos preventivos formalizados."),
        ("total_con_iva", "Numérica (COP)", "Monto total facturable con el 19% de IVA y deducción de tarifa de diagnóstico.", "NEGATIVO (-0.280): Cotizaciones con total elevado (> $1.200.000 COP) experimentan mayor tasa de rechazo."),
        ("monto_repuestos", "Numérica (COP)", "Subtotal acumulado de repuestos físicos de tiendas necesarios para el trabajo.", "NEGATIVO (-0.264): Repuestos costosos encarecen la orden. Sugiere ofrecer alternativas OEM o genéricas certificadas."),
        ("num_tiendas", "Numérica entera", "Número de tiendas de autopartes distintas requeridas para abastecer la orden.", "NEGATIVO (-0.217): Múltiples tiendas involucradas generan incertidumbre en tiempos de entrega y múltiples envíos."),
        ("monto_mano_obra", "Numérica (COP)", "Valor cobrado por el taller por concepto de trabajo técnico mecánico.", "NEGATIVO (-0.178): Mano de obra desfasada frente a los promedios del mercado causa rechazo directo."),
        ("cantidad_repuestos", "Numérica entera", "Número total de unidades de repuestos cotizados.", "Permite medir la complejidad de la reparación en conjunto con la mano de obra."),
        ("proporcion_repuestos", "Numérica (0.0 - 1.0)", "Relación entre el costo de repuestos y el costo total de la cotización.", "Mide el equilibrio entre materiales y servicio técnico."),
        ("marca_vehiculo", "Categórica (One-Hot)", "Marca del auto (Chevrolet, Renault, Mazda, Toyota, Kia, Nissan, etc.).", "Segmenta el poder adquisitivo esperado y costo relativo de repuestos por fabricante."),
        ("tipo_servicio", "Categórica (One-Hot)", "Categoría del trabajo (Frenos, Suspensión, Mantenimiento preventivo, Motor, etc.).", "Distingue servicios de alta urgencia (frenos/motor) frente a servicios electivos o preventivos.")
    ]

    for row_idx, fila in enumerate(criterios_data):
        row_cells = table.add_row().cells
        bg_color = "FFFFFF" if row_idx % 2 == 0 else COLOR_FONDO_TABLA
        for col_idx, valor in enumerate(fila):
            row_cells[col_idx].text = valor
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], 100, 100, 140, 140)
            p = row_cells[col_idx].paragraphs[0]
            if col_idx == 0:
                p.runs[0].font.bold = True
            for r in p.runs:
                r.font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # --- SECCIÓN 4: ARQUITECTURA DE INTEGRACIÓN ---
    h_sec4 = doc.add_heading("4. Arquitectura de Integración en DriveMaster", level=1)
    h_sec4.style.font.color.rgb = COLOR_PRIMARIO

    doc.add_paragraph(
        "Para garantizar alta disponibilidad y desacoplamiento, el modelo se implementó bajo una "
        "arquitectura de microservicio con fallback automático nativo:"
    )

    b_a1 = doc.add_paragraph(style='List Bullet')
    r = b_a1.add_run("Microservicio en Python (FastAPI / HTTP Server): ")
    r.bold = True
    b_a1.add_run("Ejecuta en el puerto 8000 exponiendo endpoints REST (/predict, /metricas, /) procesando la inferencia en milisegundos.")

    b_a2 = doc.add_paragraph(style='List Bullet')
    r = b_a2.add_run("Backend en Java (Spring Boot 3.4.5): ")
    r.bold = True
    b_a2.add_run("El servicio PredictorCotizacionService consulta la API de Python al previsualizar la cotización. Si Python estuviese apagado, activa un motor interno de fallback en Java que resuelve la función sigmoide con los pesos exactos (modelo_pesos.json), garantizando cero fallos.")

    b_a3 = doc.add_paragraph(style='List Bullet')
    r = b_a3.add_run("Frontend en React (Vite): ")
    r.bold = True
    b_a3.add_run("El componente AlertaPrediccionCotizacion.jsx renderiza en tiempo real el porcentaje de probabilidad, el semáforo de riesgo y las recomendaciones automáticas.")

    # --- SECCIÓN 5: MÉTRICAS ACADÉMICAS ---
    h_sec5 = doc.add_heading("5. Evaluación del Modelo y Métricas Obtenidas", level=1)
    h_sec5.style.font.color.rgb = COLOR_PRIMARIO

    doc.add_paragraph(
        "El modelo fue entrenado con un dataset estructurado de 1.600 cotizaciones sintéticas realistas y evaluado "
        "en un conjunto de prueba independiente (20% = 320 cotizaciones) mediante división estratificada:"
    )

    t_met = doc.add_table(rows=1, cols=3)
    t_met.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_met.rows[0].cells[0].text = "Métrica"
    t_met.rows[0].cells[1].text = "Resultado Obtenido"
    t_met.rows[0].cells[2].text = "Interpretación Académica"
    for c in t_met.rows[0].cells:
        set_cell_background(c, COLOR_HEADER_TABLA)
        set_cell_margins(c, 120, 120, 160, 160)
        p = c.paragraphs[0]
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(9.5)

    metricas_rows = [
        ("Sensibilidad (Recall)", "82.54 %", "Capacidad destacada para capturar más del 82 % de las cotizaciones que terminarán siendo aprobadas."),
        ("Puntuación F1 (F1-Score)", "74.11 %", "Media armónica que demuestra solidez equilibrando precisión y exhaustividad."),
        ("Área Bajo la Curva (ROC-AUC)", "0.7100", "Indica una buena capacidad del modelo para distinguir cotizaciones exitosas de cotizaciones en riesgo."),
        ("Exactitud (Accuracy)", "65.94 %", "Porcentaje de acierto global sobre el conjunto de prueba independiente."),
        ("Verdaderos Positivos (TP)", "156", "Cotizaciones aprobadas predichas con éxito."),
        ("Verdaderos Negativos (TN)", "55", "Cotizaciones rechazadas anticipadas a tiempo.")
    ]

    for idx, (m, val, desc) in enumerate(metricas_rows):
        rc = t_met.add_row().cells
        bg = "FFFFFF" if idx % 2 == 0 else COLOR_FONDO_TABLA
        for col_i, texto in enumerate([m, val, desc]):
            rc[col_i].text = texto
            set_cell_background(rc[col_i], bg)
            set_cell_margins(rc[col_i], 90, 90, 120, 120)
            p = rc[col_i].paragraphs[0]
            if col_i == 0:
                p.runs[0].font.bold = True
            if col_i == 1:
                p.runs[0].font.bold = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # --- SECCIÓN 6: CÓMO SE USA EN LA APP ---
    h_sec6 = doc.add_heading("6. Flujo de Uso en la Aplicación (Paso a Paso)", level=1)
    h_sec6.style.font.color.rgb = COLOR_PRIMARIO

    pasos = [
        ("Paso 1: Recepción de la Solicitud", "El cliente sube su auto, describe el problema y selecciona repuestos si los conoce. El taller recibe la solicitud en su bandeja de entrada."),
        ("Paso 2: Formulación de la Cotización", "El taller ingresa los trabajos de mano de obra (por horas o valor fijo)."),
        ("Paso 3: Cálculo y Evaluación Predictiva", "Al hacer clic en 'Calcular totales', el backend envía las variables al modelo de Machine Learning, el cual evalúa los 12 criterios y responde en 20 ms."),
        ("Paso 4: Interpretación de los Niveles de Riesgo", "El taller visualiza un panel con el indicador de riesgo:\n"
         "  • Riesgo Bajo (Verde, ≥ 70%): La cotización es competitiva y se recomienda enviarla.\n"
         "  • Riesgo Moderado (Ámbar, 45% - 69%): Se sugiere revisar repuestos alternativos o afinar mano de obra.\n"
         "  • Riesgo Alto (Rojo, < 45%): Alerta de alta probabilidad de rechazo con sugerencias puntuales (ej. demora en la respuesta o monto excesivo)."),
        ("Paso 5: Ajuste o Emisión Final", "El taller ajusta su propuesta según las sugerencias o la emite directamente con plena consciencia de su probabilidad de éxito.")
    ]

    for titulo, desc in pasos:
        p = doc.add_paragraph()
        r = p.add_run(f"{titulo}\n")
        r.bold = True
        r.font.color.rgb = COLOR_SECUNDARIO
        p.add_run(desc)

    # Guardar documento
    OUTPUT_PATH = os.path.join(r"c:\Users\Sebastian Ortega\Desktop\DriveMaster\Drivemaster", "Modelo_Predictivo_DriveMaster.docx")
    doc.save(OUTPUT_PATH)
    print(f"Documento Word generado exitosamente en: {OUTPUT_PATH}")

if __name__ == "__main__":
    crear_documento()

